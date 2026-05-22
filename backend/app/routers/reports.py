from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from datetime import date
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.cell.cell import MergedCell
from .. import models
from ..database import get_db
from ..auth import get_current_user
import os

_BASE_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

router = APIRouter(prefix="/reports", tags=["reports"])

HEADER_FILL = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ALT_FILL = PatternFill(start_color="EBF3FB", end_color="EBF3FB", fill_type="solid")


def style_headers(ws, headers):
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 20


@router.get("/equipment")
async def export_equipment_report(
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventario de Equipos"

    headers = ["Código", "Nombre", "Categoría", "Marca", "Modelo", "N° Serie",
               "Ubicación", "Área", "Estado", "Criticidad", "F. Compra",
               "F. Instalación", "Garantía", "Último Mant.", "Próx. Mant.", "Costo Total"]
    style_headers(ws, headers)

    equipments = db.query(models.Equipment).all()
    for i, eq in enumerate(equipments, 2):
        row_data = [
            eq.code, eq.name,
            eq.category_obj.name if eq.category_obj else "",
            eq.brand or "", eq.model or "", eq.serial_number or "",
            eq.location or "", eq.area or "",
            str(eq.status.value if eq.status else ""),
            str(eq.criticality.value if eq.criticality else ""),
            str(eq.purchase_date) if eq.purchase_date else "",
            str(eq.installation_date) if eq.installation_date else "",
            str(eq.warranty_expiry) if eq.warranty_expiry else "",
            str(eq.last_maintenance) if eq.last_maintenance else "",
            str(eq.next_maintenance) if eq.next_maintenance else "",
            eq.total_maintenance_cost or 0,
        ]
        for col, value in enumerate(row_data, 1):
            ws.cell(row=i, column=col, value=value)
            if i % 2 == 0:
                ws.cell(row=i, column=col).fill = ALT_FILL

    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=equipos_{date.today()}.xlsx"}
    )


@router.get("/work-orders")
async def export_work_orders_report(
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Órdenes de Trabajo"

    headers = ["N° OT", "Equipo", "Código", "Tipo", "Prioridad", "Estado",
               "F. Programada", "F. Inicio", "F. Fin", "Técnico",
               "Horas Labor", "Costo Labor", "Costo Mat.", "Costo Total",
               "Horas Parada", "Hallazgos", "Acciones"]
    style_headers(ws, headers)

    wos = db.query(models.WorkOrder).order_by(models.WorkOrder.created_at.desc()).all()
    for i, wo in enumerate(wos, 2):
        tech_name = wo.technician.full_name if wo.technician else ""
        eq_name = wo.equipment.name if wo.equipment else ""
        eq_code = wo.equipment.code if wo.equipment else ""
        row_data = [
            wo.order_number, eq_name, eq_code,
            str(wo.type.value if wo.type else ""),
            str(wo.priority.value if wo.priority else ""),
            str(wo.status.value if wo.status else ""),
            str(wo.scheduled_date) if wo.scheduled_date else "",
            str(wo.start_datetime) if wo.start_datetime else "",
            str(wo.end_datetime) if wo.end_datetime else "",
            tech_name,
            wo.labor_hours or 0, wo.labor_cost or 0,
            wo.materials_cost or 0, wo.total_cost or 0,
            wo.downtime_hours or 0,
            wo.findings or "", wo.actions_taken or "",
        ]
        for col, value in enumerate(row_data, 1):
            ws.cell(row=i, column=col, value=value)
            if i % 2 == 0:
                ws.cell(row=i, column=col).fill = ALT_FILL

    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=ordenes_trabajo_{date.today()}.xlsx"}
    )


@router.get("/lifecycle/{equipment_id}")
async def export_lifecycle_report(
    equipment_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    eq = db.query(models.Equipment).filter(models.Equipment.id == equipment_id).first()
    if not eq:
        from fastapi import HTTPException
        raise HTTPException(status_code=404, detail="Equipo no encontrado")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Hoja de Vida - {eq.code}"

    # Header info
    ws.merge_cells("A1:H1")
    title_cell = ws.cell(row=1, column=1, value=f"HOJA DE VIDA DEL EQUIPO: {eq.name} ({eq.code})")
    title_cell.font = Font(bold=True, size=14, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 25

    info = [
        ("Marca:", eq.brand or ""), ("Modelo:", eq.model or ""),
        ("N° Serie:", eq.serial_number or ""), ("Ubicación:", eq.location or ""),
        ("Estado:", str(eq.status.value if eq.status else "")),
        ("Criticidad:", str(eq.criticality.value if eq.criticality else "")),
        ("F. Instalación:", str(eq.installation_date) if eq.installation_date else ""),
        ("Garantía:", str(eq.warranty_expiry) if eq.warranty_expiry else ""),
    ]
    for row_offset, (label, value) in enumerate(info, 2):
        ws.cell(row=row_offset, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row_offset, column=2, value=value)

    # Embed equipment photo if available
    if eq.image_url:
        rel = eq.image_url.lstrip("/")
        img_path = os.path.join(_BASE_DIR, rel)
        if os.path.isfile(img_path):
            try:
                xl_img = XLImage(img_path)
                xl_img.width = 180
                xl_img.height = 140
                ws.add_image(xl_img, "E2")
            except Exception:
                pass  # skip if image format unsupported

    start_row = len(info) + 3
    headers = ["Fecha", "Tipo", "Descripción", "Hallazgos", "Acciones", "Técnico", "Horas", "Costo"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[start_row].height = 20

    entries = db.query(models.EquipmentLifecycle).filter(
        models.EquipmentLifecycle.equipment_id == equipment_id
    ).order_by(models.EquipmentLifecycle.date.asc()).all()

    for i, entry in enumerate(entries, start_row + 1):
        tech_name = entry.technician.full_name if entry.technician else ""
        row_data = [
            str(entry.date)[:10] if entry.date else "",
            str(entry.event_type.value if entry.event_type else ""),
            entry.description or "",
            entry.findings or "",
            entry.actions_taken or "",
            tech_name,
            entry.labor_hours or 0,
            entry.total_cost or 0,
        ]
        for col, value in enumerate(row_data, 1):
            ws.cell(row=i, column=col, value=value)
            if i % 2 == 0:
                ws.cell(row=i, column=col).fill = ALT_FILL

    for col in ws.columns:
        first = next((c for c in col if not isinstance(c, MergedCell)), None)
        if first is None:
            continue
        max_len = max((len(str(c.value or "")) for c in col if not isinstance(c, MergedCell)), default=10)
        ws.column_dimensions[first.column_letter].width = min(max_len + 2, 50)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=hoja_vida_{eq.code}_{date.today()}.xlsx"}
    )
